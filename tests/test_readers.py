"""Tests de extracción de texto por tipo de archivo."""

from __future__ import annotations

import io

import pytest

from westfield_agent_ingest.readers import extract_text, is_supported


def test_is_supported() -> None:
    assert is_supported("caso.pdf")
    assert is_supported("notas.DOCX")
    assert is_supported("guia.md")
    assert is_supported("texto.txt")
    assert not is_supported("imagen.png")
    assert not is_supported("docs.json")


def test_extract_md_y_txt_decodifican_utf8() -> None:
    assert extract_text("a.md", "# Título\ncontenido".encode()) == "# Título\ncontenido"
    assert extract_text("b.txt", b"texto plano") == "texto plano"


def test_extension_no_soportada_es_error() -> None:
    with pytest.raises(ValueError, match="no soportada"):
        extract_text("imagen.png", b"...")


def test_extract_docx_roundtrip() -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("Primer párrafo del caso.")
    document.add_paragraph("Segundo párrafo con detalle.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("caso.docx", buffer.getvalue())
    assert "Primer párrafo del caso." in text
    assert "Segundo párrafo con detalle." in text
